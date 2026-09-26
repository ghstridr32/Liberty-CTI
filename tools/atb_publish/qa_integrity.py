"""
qa_integrity.py -- Automated content-integrity comparison between the parsed
canonical ATB source and the generated DOCX/PDF companion files.

Fails (returns ok=False) if:
  - extracted text from the companion is far shorter than the source (major
    content likely dropped during generation), or
  - a heading present in the source cannot be found anywhere in the companion
    text (a whole section likely went missing).

This is deliberately a coarse, high-recall check: it is meant to catch gross
omissions (a dropped section, a failed table render, truncation), not to
verify exact wording. It does not, and cannot, judge analytic accuracy.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader

sys.path.insert(0, str(Path(__file__).parent))
from parse_atb import parse_issue, runs_to_plain  # noqa: E402

MIN_RATIO = 0.55


def _normalize(s: str) -> str:
    return re.sub(r"\s+", " ", s or "").strip().lower()


def extract_docx_text(path: str | Path) -> str:
    doc = DocxDocument(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def extract_pdf_text(path: str | Path) -> str:
    reader = PdfReader(str(path))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _source_headings(issue: dict) -> list[str]:
    return [b["text"] for b in issue["blocks"] if b["type"] == "heading" and len(b["text"]) > 3]


def check_companion(issue: dict, companion_text: str, label: str) -> dict:
    errors = []
    warnings = []

    source_len = issue.get("_source_text_len", 0)
    companion_len = len(_normalize(companion_text))
    ratio = (companion_len / source_len) if source_len else 0

    if ratio < MIN_RATIO:
        errors.append(
            f"{label}: extracted text is only {ratio:.0%} of source length "
            f"({companion_len} vs {source_len} chars) -- likely missing content."
        )
    elif ratio < 0.8:
        warnings.append(f"{label}: extracted text is {ratio:.0%} of source length -- review recommended.")

    norm_companion = _normalize(companion_text)
    missing_headings = []
    for heading in _source_headings(issue):
        if _normalize(heading) not in norm_companion:
            missing_headings.append(heading)
    if missing_headings:
        errors.append(f"{label}: {len(missing_headings)} source heading(s) not found in companion: {missing_headings}")

    return {"ok": len(errors) == 0, "errors": errors, "warnings": warnings, "ratio": ratio}


def check_issue_companions(issue: dict, docx_path: str | Path, pdf_path: str | Path) -> dict:
    docx_text = extract_docx_text(docx_path)
    pdf_text = extract_pdf_text(pdf_path)

    docx_result = check_companion(issue, docx_text, "DOCX")
    pdf_result = check_companion(issue, pdf_text, "PDF")

    ok = docx_result["ok"] and pdf_result["ok"]
    return {
        "ok": ok,
        "docx": docx_result,
        "pdf": pdf_result,
        "errors": docx_result["errors"] + pdf_result["errors"],
        "warnings": docx_result["warnings"] + pdf_result["warnings"],
    }


def main():
    if len(sys.argv) < 4:
        print("usage: qa_integrity.py <issue.html> <out.docx> <out.pdf>", file=sys.stderr)
        sys.exit(1)
    issue = parse_issue(sys.argv[1])
    result = check_issue_companions(issue, sys.argv[2], sys.argv[3])
    for e in result["errors"]:
        print(f"ERROR: {e}")
    for w in result["warnings"]:
        print(f"WARNING: {w}")
    print(f"DOCX ratio: {result['docx']['ratio']:.2%}  PDF ratio: {result['pdf']['ratio']:.2%}")
    if result["ok"]:
        print("INTEGRITY QA PASSED")
    else:
        print("ATB COMPANION BUILD FAILED -- INTEGRITY QA FAILED")
    sys.exit(0 if result["ok"] else 3)


if __name__ == "__main__":
    main()
