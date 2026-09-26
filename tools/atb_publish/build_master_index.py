"""
build_master_index.py -- Build the yearly Alamo Threat Brief master index
(DOCX + PDF) from the meta.json bookkeeping files written by archive_issue.py.

  ATB Archive/<year>/Alamo_Threat_Brief_<year>_Master_Index.docx
  ATB Archive/<year>/Alamo_Threat_Brief_<year>_Master_Index.pdf

Usage:
    python build_master_index.py 2026
"""
from __future__ import annotations

import json
import sys
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import BaseDocTemplate, Frame, PageTemplate, Paragraph, Spacer, Table, TableStyle

sys.path.insert(0, str(Path(__file__).parent))
from build_docx import GOLD_DARK, GRAY, NAVY, _set_cell_bg, _set_cell_border  # noqa: E402

REPO_ROOT = Path(__file__).resolve().parents[2]
ARCHIVE_ROOT = REPO_ROOT / "ATB Archive"


def collect_year_rows(year: int) -> list[dict]:
    year_dir = ARCHIVE_ROOT / str(year)
    rows = []
    if not year_dir.exists():
        return rows
    for issue_dir in sorted(year_dir.iterdir()):
        meta_path = issue_dir / "meta.json"
        if not meta_path.exists():
            continue
        try:
            m = json.loads(meta_path.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            continue
        rows.append(m)
    rows.sort(key=lambda r: r.get("issue_date") or "")
    return rows


def build_docx_index(year: int, rows: list[dict], out_path: Path):
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Cm(29.7), Cm(21.0)
    section.left_margin = section.right_margin = Cm(1.5)
    section.top_margin = section.bottom_margin = Cm(1.5)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"THE ALAMO THREAT BRIEF — {year} MASTER INDEX")
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = NAVY

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run("Liberty CTI  |  Archival record of every published issue and its companion files")
    sr.italic = True
    sr.font.size = Pt(10)
    sr.font.color.rgb = GRAY

    stats = doc.add_paragraph()
    stats.alignment = WD_ALIGN_PARAGRAPH.CENTER
    total = len(rows)
    first_date = rows[0]["issue_date"] if rows else "—"
    last_date = rows[-1]["issue_date"] if rows else "—"
    gen_date = date.today().isoformat()
    stat_run = stats.add_run(
        f"Total issues: {total}   |   First issue: {first_date}   |   "
        f"Latest issue: {last_date}   |   Archive generated: {gen_date}"
    )
    stat_run.font.size = Pt(9.5)
    stat_run.font.color.rgb = GOLD_DARK
    stat_run.bold = True

    doc.add_paragraph()

    headers = ["Issue", "Coverage", "Publication", "Title / Dominant Judgment", "DOCX", "PDF", "Source"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [2.6, 4.5, 2.6, 9.5, 2.8, 2.8, 3.2]
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        _set_cell_bg(hdr_cells[i], "142233")
        hdr_cells[i].width = Cm(widths[i])
        rp = hdr_cells[i].paragraphs[0]
        rr = rp.add_run(h.upper())
        rr.bold = True
        rr.font.size = Pt(9)
        rr.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for m in rows:
        cells = table.add_row().cells
        values = [
            m.get("issue_number", ""),
            m.get("coverage_period") or "",
            m.get("issue_date", ""),
            m.get("dominant_theme") or (m.get("title_raw") or "")[:90],
            m.get("docx_file", ""),
            m.get("pdf_file", ""),
            m.get("html_file", ""),
        ]
        for i, v in enumerate(values):
            _set_cell_border(cells[i], "CCCCCC")
            cells[i].width = Cm(widths[i])
            cp = cells[i].paragraphs[0]
            cr = cp.add_run(str(v))
            cr.font.size = Pt(8.5)
            cr.font.color.rgb = NAVY

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))


def build_pdf_index(year: int, rows: list[dict], out_path: Path):
    page = landscape(LETTER)
    margin = 1.3 * cm
    doc = BaseDocTemplate(str(out_path), pagesize=page,
                           leftMargin=margin, rightMargin=margin, topMargin=margin, bottomMargin=margin,
                           title=f"Alamo Threat Brief {year} Master Index", author="Liberty CTI")
    frame = Frame(margin, margin, page[0] - 2 * margin, page[1] - 2 * margin, id="main")
    doc.addPageTemplates([PageTemplate(id="idx", frames=[frame])])

    title_style = ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=17, textColor=colors.HexColor("#142233"), alignment=1, spaceAfter=4)
    sub_style = ParagraphStyle("s", fontName="Helvetica-Oblique", fontSize=9.5, textColor=colors.HexColor("#5C5C5C"), alignment=1, spaceAfter=4)
    stat_style = ParagraphStyle("st", fontName="Helvetica-Bold", fontSize=9, textColor=colors.HexColor("#7A5A1E"), alignment=1, spaceAfter=12)
    cell_style = ParagraphStyle("c", fontName="Helvetica", fontSize=7.5, textColor=colors.HexColor("#142233"), leading=9.5)
    hdr_style = ParagraphStyle("h", fontName="Helvetica-Bold", fontSize=8, textColor=colors.white)

    total = len(rows)
    first_date = rows[0]["issue_date"] if rows else "—"
    last_date = rows[-1]["issue_date"] if rows else "—"
    gen_date = date.today().isoformat()

    story = [
        Paragraph(f"THE ALAMO THREAT BRIEF — {year} MASTER INDEX", title_style),
        Paragraph("Liberty CTI  |  Archival record of every published issue and its companion files", sub_style),
        Paragraph(f"Total issues: {total} &nbsp;|&nbsp; First issue: {first_date} &nbsp;|&nbsp; "
                  f"Latest issue: {last_date} &nbsp;|&nbsp; Archive generated: {gen_date}", stat_style),
    ]

    headers = ["Issue", "Coverage", "Pub.", "Title / Dominant Judgment", "DOCX", "PDF", "Source"]
    data = [[Paragraph(h.upper(), hdr_style) for h in headers]]
    for m in rows:
        data.append([
            Paragraph(m.get("issue_number", ""), cell_style),
            Paragraph(m.get("coverage_period") or "", cell_style),
            Paragraph(m.get("issue_date", ""), cell_style),
            Paragraph((m.get("dominant_theme") or (m.get("title_raw") or ""))[:110], cell_style),
            Paragraph(m.get("docx_file", ""), cell_style),
            Paragraph(m.get("pdf_file", ""), cell_style),
            Paragraph(m.get("html_file", ""), cell_style),
        ])

    col_widths = [2.4 * cm, 4.0 * cm, 2.1 * cm, 9.0 * cm, 4.4 * cm, 4.4 * cm, 3.2 * cm]
    tbl = Table(data, colWidths=col_widths, repeatRows=1)
    tbl.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#142233")),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CCCCCC")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F1E8")]),
    ]))
    story.append(tbl)
    doc.build(story)


def build_master_index(year: int) -> dict:
    rows = collect_year_rows(year)
    out_dir = ARCHIVE_ROOT / str(year)
    docx_path = out_dir / f"Alamo_Threat_Brief_{year}_Master_Index.docx"
    pdf_path = out_dir / f"Alamo_Threat_Brief_{year}_Master_Index.pdf"
    build_docx_index(year, rows, docx_path)
    build_pdf_index(year, rows, pdf_path)
    return {"year": year, "issue_count": len(rows), "docx": str(docx_path), "pdf": str(pdf_path)}


def main():
    if len(sys.argv) < 2:
        print("usage: build_master_index.py <year>", file=sys.stderr)
        sys.exit(1)
    result = build_master_index(int(sys.argv[1]))
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
